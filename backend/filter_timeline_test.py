"""
filter_timeline_test.py
-----------------------
READ-ONLY experiment: replay our 7-step entry filter day by day over the last
15 days for 10 real active coins, in the style of the STONK / PONS timeline
workbook. Nothing is written to the database or changed in the site / filter /
drop_logic.

Run (from the backend/ folder, with backend/.env present):
    python filter_timeline_test.py
Output: backend/filter_timeline_15d.xlsx

Window and coins are FIXED (Sep 24 -> Oct 8, the same 10 coins) so a re-run
reproduces the same table. The last row (Oct 8) is the "current" row.

Per coin, per day D, with the history that actually exists:
    Price      Codex getBars 1D close for D (shown only, not a filter step)
    Age        days since the first DexScreener pool (else the CMC listing date
               cached in age_lookups)                         >= 10 days
    Market Cap CoinGecko market_chart (daily); if CoinGecko has none, Codex
               close x today's circulating supply              $1M - $400M
    Volume     CoinGecko daily volume, the 5 complete days before D (exactly
               like the live check): day5 >= day1 x 1.35; no data -> FAIL
    Liquidity  Codex daily liquidity at D vs ~30 days before: down > 5% ->
               FAIL, under $5K -> FAIL (today also uses DexScreener's
               reading, the higher one counts); no data -> pass
    Revenue    DefiLlama by contract: tracked with no fees -> FAIL; the value
               shown is DefiLlama's daily revenue for D when it has one
    Honeypot   GoPlus answers for NOW only -> checked on the last row, "—" on
               past days (unknown passes, same as the live filter)
    Holders    shown only (not one of the entry steps); our own daily
               snapshots, which start Oct 5 -> "—" before that
Step 2 (meme / stablecoin / tokenized / wrapped) is today's classification -
all 10 coins are active, so they passed it.
"""

import time
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from age_check import CHAIN_TO_CG_PLATFORM
from database import SessionLocal
from discovery import MAX_MCAP, MIN_MCAP
from entry_zones_test import CODEX_KEY, CODEX_URL
from liquidity_check import CHAIN_TO_NETWORK, MAX_DROP_PCT, MIN_LIQUIDITY_USD
from models import AgeLookup, Coin, DailySnapshot
from revenue_check import check_revenue
from revenue_detail import FEES_URL, _get, resolve_slugs
from security_check import check_token_security
from volume_check import CG_KEY, CG_MC_URL, VOLUME_MIN_RISE

DAYS = 15
MIN_AGE = 10
LIQ_WINDOW = 30
OUT_FILE = Path(__file__).resolve().parent / "filter_timeline_15d.xlsx"
END = date(2026, 10, 8)                     # last day of the fixed window
DATES = [END - timedelta(days=i) for i in range(DAYS - 1, -1, -1)]
# The 10 coins of the first run (Aethir, Beam, SAND, BENQI, OPEN, Bless,
# MAGIC, Kryll, Smooth Love Potion, Peapods) - pinned by id.
COIN_IDS = [813, 814, 815, 816, 819, 821, 822, 823, 826, 827]

BARS_QUERY = """
query ($symbol: String!, $from: Int!, $to: Int!) {
  getBars(symbol: $symbol, from: $from, to: $to, resolution: "1D",
          removeLeadingNullValues: true) { t c volume liquidity }
}"""


def _day(ts):
    return datetime.fromtimestamp(ts, tz=timezone.utc).date()


def codex_bars(contract, chain):
    """{date: {"c", "volume", "liquidity"}} from Codex getBars 1D (~50 days)."""
    net = CHAIN_TO_NETWORK.get((chain or "").lower())
    if not net:
        return {}
    now = int(time.time())
    frm = int(datetime.combine(DATES[0], datetime.min.time(), timezone.utc).timestamp())         - (LIQ_WINDOW + 5) * 86400
    try:
        body = requests.post(CODEX_URL, json={"query": BARS_QUERY, "variables": {
            "symbol": f"{contract}:{net}", "from": frm, "to": now}},
            headers={"Authorization": CODEX_KEY}, timeout=30).json()
    except (requests.exceptions.RequestException, ValueError):
        return {}
    b = ((body or {}).get("data") or {}).get("getBars") or {}
    out = {}
    for i, t in enumerate(b.get("t") or []):
        def val(k):
            arr = b.get(k) or []
            try:
                return float(arr[i]) if arr[i] is not None else None
            except (IndexError, TypeError, ValueError):
                return None
        out[_day(t)] = {"c": val("c"), "volume": val("volume"), "liquidity": val("liquidity")}
    return out


def coingecko_chart(contract, chain):
    """({date: mcap}, {date: volume}) - last reading of each UTC day, the same
    collapse volume_check uses. Empty dicts if CoinGecko has nothing."""
    platform = CHAIN_TO_CG_PLATFORM.get((chain or "").lower())
    if not platform or not CG_KEY:
        return {}, {}
    time.sleep(2.5)
    try:
        r = requests.get(CG_MC_URL.format(platform=platform, contract=contract),
                         params={"vs_currency": "usd",
                                 "days": (datetime.now(timezone.utc).date() - DATES[0]).days + 8},
                         headers={"x-cg-demo-api-key": CG_KEY}, timeout=30)
        data = r.json() if r.status_code == 200 else {}
    except (requests.exceptions.RequestException, ValueError):
        data = {}
    mcap, vol = {}, {}
    for ts, v in data.get("market_caps") or []:
        if v:
            mcap[_day(ts / 1000)] = v
    for ts, v in data.get("total_volumes") or []:
        vol[_day(ts / 1000)] = v
    return mcap, vol


def dexscreener(contract, chain):
    """(first pool creation datetime on this chain, current max liquidity)."""
    try:
        pairs = requests.get(f"https://api.dexscreener.com/latest/dex/tokens/{contract}",
                             timeout=20).json().get("pairs") or []
    except (requests.exceptions.RequestException, ValueError):
        return None, None
    pairs = [p for p in pairs if (p.get("chainId") or "").lower() == chain.lower()] or pairs
    created = [p["pairCreatedAt"] for p in pairs if p.get("pairCreatedAt")]
    liqs = [(p.get("liquidity") or {}).get("usd") or 0 for p in pairs]
    first = datetime.fromtimestamp(min(created) / 1000, tz=timezone.utc) if created else None
    return first, (max(liqs) if liqs else None)


def defillama_daily(contract):
    """{date: daily revenue} from DefiLlama (first slug that has a chart)."""
    for slug in resolve_slugs(contract):
        fees, ok = _get(FEES_URL.format(slug=slug), {"dataType": "dailyRevenue"})
        chart = (fees or {}).get("totalDataChart") or []
        if chart:
            return {_day(t): float(v or 0) for t, v in chart}
    return {}


def money(v):
    if v is None:
        return "—"
    for div, suf in ((1e9, "B"), (1e6, "M"), (1e3, "K")):
        if abs(v) >= div:
            return f"${v / div:.1f}{suf}".replace(".0" + suf, suf)
    return f"${v:,.0f}"


def price(v):
    """Close price with enough significant digits for sub-cent tokens."""
    if v is None:
        return "—"
    if v >= 1:
        return f"${v:,.4f}"
    return f"${v:.4g}" if v >= 0.0001 else f"${v:.3e}"


def evaluate(coin, session):
    contract, chain = coin.contract, coin.chain
    bars = codex_bars(contract, chain)
    cg_mcap, cg_vol = coingecko_chart(contract, chain)
    first_pool, dex_liq = dexscreener(contract, chain)
    born, born_src = first_pool, "DexScreener"
    if born is None:
        row = session.query(AgeLookup).filter_by(chain=chain.lower(), contract=contract.lower()).first()
        if row and row.listed_at:
            born, born_src = row.listed_at.replace(tzinfo=timezone.utc), "CMC listing"
    rev = check_revenue(coin.name, None, contract, chain)
    rev_daily = defillama_daily(contract) if rev["verdict"] == "HAS_REVENUE" else {}
    sec = check_token_security(contract, chain)
    holders = {s.date: s.holders for s in session.query(DailySnapshot)
               .filter(DailySnapshot.coin_id == coin.id, DailySnapshot.holders.isnot(None))}
    liq_series = sorted((d, b["liquidity"]) for d, b in bars.items() if b["liquidity"] is not None)

    rows = []
    for d in DATES:
        fails, notes = [], []
        # Step 3 - age
        age = (d - born.date()).days if born else None
        if age is None:
            fails.append("age unknown")
        elif age < MIN_AGE:
            fails.append(f"age {age}d < {MIN_AGE}")
        # Step 1 - market cap
        mcap, mcap_note = cg_mcap.get(d), ""
        if mcap is None and bars.get(d, {}).get("c") and coin.circulating_supply:
            mcap, mcap_note = bars[d]["c"] * coin.circulating_supply, " (est.)"
        if mcap is None:
            notes.append("no market cap for this day")
        elif not MIN_MCAP <= mcap <= MAX_MCAP:
            fails.append(f"MCap {money(mcap)} outside $1M–$400M")
        # Step 4 - revenue
        if rev["verdict"] == "NO_REVENUE":
            fails.append("DefiLlama: no fees")
            rev_txt = "No fees"
        elif rev["verdict"] == "HAS_REVENUE":
            rev_txt = f"{money(rev_daily[d])}/day" if d in rev_daily else "Yes"
        else:
            rev_txt = "Not tracked"
        # Step 5 - volume (5 complete days before D)
        win = [cg_vol.get(d - timedelta(days=k)) for k in range(5, 0, -1)]
        if None in win or not win[0]:
            vol_txt = "—"
            fails.append("no CoinGecko volume history")
        else:
            chg = (win[-1] / win[0] - 1) * 100
            vol_txt = f"{chg:+.0f}% ({money(win[0])}→{money(win[-1])})"
            if win[-1] < win[0] * VOLUME_MIN_RISE:
                fails.append(f"volume {chg:+.0f}% (needs +35%)")
        # Step 6 - honeypot (today only)
        if d == END:
            hp = {"safe": "Safe", "unsafe": "Unsafe"}.get(sec["verdict"], "Unknown")
            if sec["verdict"] == "unsafe":
                fails.append("honeypot / can't sell")
        else:
            hp = "—"
        # Step 7 - liquidity
        liq = bars.get(d, {}).get("liquidity")
        if liq is None:
            liq_txt = "—"
            notes.append("no Codex liquidity (passes as unknown)")
        else:
            base = next((v for dd, v in liq_series if d - timedelta(days=LIQ_WINDOW) <= dd <= d), None)
            pct = (liq / base - 1) * 100 if base else None
            liq_txt = money(liq) + (f" ({pct:+.0f}%)" if pct is not None else "")
            best = max(liq, dex_liq or 0) if d == END else liq
            if pct is not None and pct < -MAX_DROP_PCT:
                fails.append(f"liquidity {pct:+.0f}% vs 30d")
            if best < MIN_LIQUIDITY_USD:
                fails.append(f"liquidity {money(best)} < $5K")
        h = holders.get(d)
        close = bars.get(d, {}).get("c")
        rows.append({
            "date": d.isoformat(), "price": price(close), "age": age if age is not None else "—",
            "mcap": money(mcap) + mcap_note, "vol": vol_txt,
            "holders": f"{h:,}" if h else "—", "liq": liq_txt, "rev": rev_txt, "hp": hp,
            "pass": not fails, "fails": fails,
            "note": "; ".join(fails + notes) if (fails or notes) else "All steps pass",
        })
    meta = {"born": born, "born_src": born_src, "rev": rev["verdict"], "sec": sec["verdict"]}
    return rows, meta


# ---------------------------------------------------------------- workbook
NAVY, PALE = "1F3A5F", "D9E1F2"
F_HEAD = PatternFill("solid", start_color=NAVY)
F_SUB = PatternFill("solid", start_color=PALE)
F_YES = PatternFill("solid", start_color="C6EFCE")
F_NO = PatternFill("solid", start_color="FFC7CE")
F_HI = PatternFill("solid", start_color="FFEB9C")
BASE = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
WHITE = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TITLE = Font(name="Arial", size=14, bold=True, color=NAVY)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
HEADERS = ["Date", "Price (USD)", "Age (days)", "Market Cap", "Volume (5d)", "Holders", "Liquidity",
           "Fees/Revenue", "Honeypot", "Passes Filter?", "Note"]
WIDTHS = [12, 13, 10, 14, 24, 11, 16, 13, 10, 13, 48]


def sheet_name(name, used):
    base = "".join(ch for ch in name if ch not in '[]:*?/\\')[:28] or "Coin"
    n, i = base, 2
    while n in used:
        n, i = f"{base[:25]} {i}", i + 1
    used.add(n)
    return n


def coin_sheet(wb, title, coin, rows, meta, used):
    ws = wb.create_sheet(sheet_name(coin.name, used))
    ws["A1"] = f"{coin.name} — Filter Application Timeline (last {DAYS} days, {coin.chain})"
    ws["A1"].font = TITLE
    ws.merge_cells("A1:K1")
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.fill, cell.font, cell.alignment = F_HEAD, WHITE, CENTER
    ws.row_dimensions[3].height = 27.75
    first = next((i for i, r in enumerate(rows) if r["pass"]), None)
    for i, r in enumerate(rows):
        vals = [r["date"], r["price"], r["age"], r["mcap"], r["vol"], r["holders"], r["liq"], r["rev"],
                r["hp"], "YES" if r["pass"] else "NO",
                ("First pass in this window — " + r["note"]) if i == first else r["note"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=4 + i, column=c, value=v)
            cell.font, cell.alignment = BASE, (LEFT if c == 11 else CENTER)
            if i == first:
                cell.fill = F_HI
        pc = ws.cell(row=4 + i, column=10)
        pc.fill, pc.font = (F_YES if r["pass"] else F_NO), BOLD
        if i == first:
            ws.cell(row=4 + i, column=11).font = BOLD
    for c, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[chr(64 + c)].width = w
    end = 4 + len(rows) + 1
    passed = sum(r["pass"] for r in rows)
    ws.cell(row=end, column=1, value="RESULT:").font = BOLD
    born = meta["born"].date().isoformat() if meta["born"] else "unknown"
    ws.cell(row=end, column=3, value=(
        f"Passed {passed} of {DAYS} days  •  Created {born} ({meta['born_src']})  •  "
        f"Revenue: {meta['rev']}  •  Honeypot (checked {datetime.now(timezone.utc).date()}): {meta['sec']}")).font = BASE
    ws.merge_cells(start_row=end, start_column=1, end_row=end, end_column=2)
    ws.merge_cells(start_row=end, start_column=3, end_row=end, end_column=11)


def summary_sheet(ws, results):
    ws.title = "Summary"
    ws["A1"] = "Filter Summary & Results"
    ws["A1"].font = TITLE

    def band(row, text, cols=4):
        c = ws.cell(row=row, column=1, value=text)
        c.fill, c.font = F_HEAD, WHITE
        for k in range(2, cols + 1):
            ws.cell(row=row, column=k).fill = F_HEAD

    def sub(row, labels):
        for k, t in enumerate(labels, 1):
            c = ws.cell(row=row, column=k, value=t)
            c.fill, c.font, c.alignment = F_SUB, BOLD, CENTER

    band(3, "THE FILTER (our real 7 steps, replayed per day)")
    sub(4, ["#", "Criterion", "Threshold", "Source (historical)"])
    crit = [
        (1, "Market Cap", "$1M – $400M", "CoinGecko daily (else Codex price × supply, marked est.)"),
        (2, "Not meme / stable / tokenized / wrapped", "Excluded categories", "Today's classification (all 10 passed)"),
        (3, "Age", f">= {MIN_AGE} days", "DexScreener first pool (else CMC listing)"),
        (4, "Fees OR Revenue", "Not 'tracked with no fees'", "DefiLlama by contract; daily revenue shown"),
        (5, "Volume (5d)", "day5 >= day1 × 1.35", "CoinGecko daily volume, 5 complete days before the date"),
        (6, "Honeypot", "Safe (can sell)", "GoPlus — current only, so checked on the last row"),
        (7, "Liquidity", "Not down > 5% vs 30d, and >= $5K", "Codex daily liquidity (+ DexScreener today)"),
    ]
    for i, row in enumerate(crit):
        for k, v in enumerate(row, 1):
            c = ws.cell(row=5 + i, column=k, value=v)
            c.font, c.alignment = BASE, (CENTER if k == 1 else LEFT)

    r0 = 5 + len(crit) + 1
    band(r0, "RESULTS", 5)
    sub(r0 + 1, ["Coin", "Chain", "Days Passed", "Main reason for NO days", "Last day"])
    for i, (coin, rows, _) in enumerate(results):
        passed = sum(x["pass"] for x in rows)
        reasons = Counter(f.split(" ")[0] if not f.startswith("no ") else f
                          for x in rows for f in x["fails"])
        top = ", ".join(f"{k} ({n}d)" for k, n in reasons.most_common(2)) or "—"
        vals = [coin.name, coin.chain, f"{passed} / {DAYS}", top,
                "YES" if rows[-1]["pass"] else "NO"]
        for k, v in enumerate(vals, 1):
            c = ws.cell(row=r0 + 2 + i, column=k, value=v)
            c.font, c.alignment = BASE, (LEFT if k in (1, 4) else CENTER)
        ws.cell(row=r0 + 2 + i, column=3).fill = F_YES if passed else F_NO
        ws.cell(row=r0 + 2 + i, column=5).fill = F_YES if rows[-1]["pass"] else F_NO

    n = r0 + 2 + len(results) + 1
    ws.cell(row=n, column=1, value="DATA NOTES:").font = BOLD
    for k, t in enumerate([
        f"1. Window: {DATES[0]} → {DATES[-1]} (fixed). Price (USD) = Codex daily close, shown only.",
        "2. Holders: we only keep daily holder counts since Oct 5 — earlier days show '—'. Holders is not an entry step.",
        "3. Honeypot: GoPlus only answers for now — past days show '—' and pass (unknown passes, as in the live filter).",
        "4. Fees/Revenue: '$X/day' = DefiLlama daily revenue that day; 'Not tracked' = DefiLlama doesn't list the contract (passes).",
        "5. Volume (5d): % change from the oldest to the newest of the 5 complete days before that date.",
        "6. Liquidity: Codex value that day and its % change vs ~30 days earlier.",
        "7. Read-only test: nothing in the database, site, filter or drop_logic was changed.",
    ], 1):
        ws.cell(row=n + k, column=1, value=t).font = BASE
    for c, w in zip("ABCDE", [30, 12, 14, 48, 10]):
        ws.column_dimensions[c].width = w


def main():
    session = SessionLocal()
    try:
        results = []
        for cid in COIN_IDS:
            coin = session.get(Coin, cid)
            print(f"  {coin.name} ({coin.chain}) ...", flush=True)
            rows, meta = evaluate(coin, session)
            results.append((coin, rows, meta))
    finally:
        session.close()

    wb = Workbook()
    summary_sheet(wb.active, results)
    used = {"Summary"}
    for coin, rows, meta in results:
        coin_sheet(wb, None, coin, rows, meta, used)
    wb.save(OUT_FILE)
    print(f"\nWrote {OUT_FILE}\n")
    for coin, rows, meta in results:
        print(f"  {coin.name[:28]:<28} {coin.chain:<10} passed {sum(r['pass'] for r in rows):>2}/{DAYS}  "
              f"| {END}: {'YES' if rows[-1]['pass'] else 'NO '} {rows[-1]['note'][:70]}")


if __name__ == "__main__":
    main()
